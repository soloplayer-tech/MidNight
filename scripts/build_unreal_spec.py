#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Unframed 언리얼 전용 명세서(데이터 계약서) xlsx 생성.
SSOT는 docs/명세서/Unframed_API명세서.xlsx (백엔드). 이 파일은 언리얼팀 시점으로 재포장.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY = "1B2A41"
GOLD = "D4A85A"
PANEL = "2C3A52"
LIGHT = "F2EFE6"
INK = "1A1A1A"

title_font = Font(name="Malgun Gothic", size=16, bold=True, color="FFFFFF")
sub_font   = Font(name="Malgun Gothic", size=10, color=GOLD)
sec_font   = Font(name="Malgun Gothic", size=11, bold=True, color=NAVY)
hdr_font   = Font(name="Malgun Gothic", size=10, bold=True, color="FFFFFF")
cell_font  = Font(name="Malgun Gothic", size=10, color=INK)
mono_font  = Font(name="Consolas", size=9, color=INK)

navy_fill  = PatternFill("solid", fgColor=NAVY)
gold_fill  = PatternFill("solid", fgColor=GOLD)
panel_fill = PatternFill("solid", fgColor=PANEL)
light_fill = PatternFill("solid", fgColor=LIGHT)

wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)
center = Alignment(horizontal="center", vertical="center")
thin = Side(style="thin", color="C9C2B0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_table(ws, start_row, header, rows, widths, mono_cols=()):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    r = start_row
    for c, h in enumerate(header, 1):
        cell = ws.cell(r, c, h)
        cell.font = hdr_font
        cell.fill = navy_fill
        cell.alignment = center
        cell.border = border
    r += 1
    for row in rows:
        for c, val in enumerate(row, 1):
            cell = ws.cell(r, c, val)
            cell.font = mono_font if (c in mono_cols) else cell_font
            cell.alignment = wrap
            cell.border = border
            cell.fill = light_fill if (r % 2 == 0) else PatternFill()
        r += 1
    return r


def banner(ws, title, sub):
    ws.merge_cells("A1:F1"); ws.merge_cells("A2:F2")
    a = ws["A1"]; a.value = title; a.font = title_font; a.fill = navy_fill
    a.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30
    b = ws["A2"]; b.value = sub; b.font = sub_font; b.fill = navy_fill
    b.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 18


def section(ws, row, text):
    ws.cell(row, 1, text).font = sec_font
    ws.cell(row, 1).fill = gold_fill
    for c in range(2, 7):
        ws.cell(row, c).fill = gold_fill
    return row + 1


wb = Workbook()

# ---------------- README ----------------
ws = wb.active; ws.title = "README"
banner(ws, "UNFRAMED — 언리얼 전용 명세서 (데이터 계약서)",
       "12팀 · 미드나잇 캠프 · 언리얼(박재현·임효택)용 · v0.1 (초안)")
r = 4
r = section(ws, r, "■ 이 문서가 뭔가요")
notes = [
    ("언리얼팀은 API 통신 자체보다 '받는 JSON이 어떤 모양인가'에 집중하면 됩니다.",),
    ("이 문서 = 언리얼이 읽을/받을 데이터의 모양 + 언리얼 구조체(USTRUCT) 매핑.",),
    ("백엔드용 통신 명세(HTTP/요청/에러)는 별도 파일: Unframed_API명세서.xlsx.",),
]
for n in notes:
    ws.cell(r, 1, n[0]).font = cell_font; ws.merge_cells(f"A{r}:F{r}"); ws.cell(r,1).alignment = wrap; r += 1
r += 1
r = section(ws, r, "■ 절대 규칙 (협업의 핵심)")
rules = [
    "백엔드가 주는 JSON == 언리얼이 읽는 JSON. 필드명 한 글자도 다르면 안 됩니다.",
    "SSOT(단일 진실)는 Unframed_API명세서.xlsx. 충돌 시 그 파일이 우선.",
    "지금은 로컬 JSON 파일(docs/data/*.json)을 읽고, 백엔드 완성 후 '파일 읽던 자리를 서버 호출로 교체'만 합니다(게임 로직 무변경).",
]
for n in rules:
    ws.cell(r, 1, "• " + n).font = cell_font; ws.merge_cells(f"A{r}:F{r}"); ws.cell(r,1).alignment = wrap; r += 1
r += 1
r = section(ws, r, "■ 게임 한 사이클 (언리얼 시점)")
flow = "세션시작 → 탐색 → 포획(찰칵) → 명화 설명/대화(TTS) → quizReady=true → 퀴즈 표시 → 답 제출(봉인) → 정답이면 스탬프 획득 / 오답이면 재도전 → 3종 완료 시 교실 복귀"
ws.cell(r, 1, flow).font = cell_font; ws.merge_cells(f"A{r}:F{r}"); ws.cell(r,1).alignment = wrap; ws.row_dimensions[r].height = 45; r += 2
r = section(ws, r, "■ 시트 구성")
sheets = [
    ("게임사이클", "단계별로 언리얼이 어떤 데이터를 받는지"),
    ("데이터구조", "Painting 모델 + 각 응답 JSON 필드"),
    ("USTRUCT매핑", "JSON 필드 ↔ 언리얼 C++ 구조체 매핑"),
    ("로컬파일&전환", "로컬 JSON 파일 목록 + 서버 전환 가이드 + 에셋 매핑"),
]
r = style_table(ws, r, ["시트", "내용"], sheets, [22, 60])
ws.sheet_view.showGridLines = False

# ---------------- 게임사이클 ----------------
ws = wb.create_sheet("게임사이클")
banner(ws, "게임 사이클 — 단계별 받는 데이터", "각 단계에서 언리얼이 호출/읽고, 받는 핵심 필드")
rows = [
    ("1", "세션 시작", "POST /sessions", "sessionId, targets[], totalStamps",
     "이번 판 봉인 목표 명화 목록·스탬프 수 받음"),
    ("2", "포획", "POST /captures", "captured, paintingId, greeting, greetingTtsUrl",
     "카메라로 찍으면 호출. greeting을 TTS로 재생"),
    ("3", "대화", "POST /paintings/{paintingId}/chat", "reply, ttsUrl, turnCount, quizReady",
     "유저 질문 전송→답변+음성. quizReady=true면 퀴즈 버튼 활성화"),
    ("4", "퀴즈 생성", "POST /paintings/{paintingId}/quiz", "quizId, question, choices[]",
     "대화 맥락 기반 객관식 1개. answerIndex는 클라에 안 옴"),
    ("5", "봉인(제출)", "POST /paintings/{paintingId}/seal", "correct, sealed, stampCount, totalStamps, feedback",
     "선택한 보기 제출. sealed=true면 봉인 연출+스탬프"),
    ("6", "도감 조회", "GET /players/{playerId}/collection", "sealedPaintings[], sealedCount",
     "재방문 시 이미 봉인한 명화 복원"),
    ("7", "진척/통계", "GET /players/{playerId}/progress", "completionRate, quizAccuracy, totalAttempts, playTimeSec",
     "발표 시연용 통계 화면"),
    ("8", "세션 종료", "POST /sessions/{sessionId}/complete", "completed, sealedCount, playTimeSec, summaryText",
     "3종 완료 시. summaryText로 교실 복귀 연출"),
]
style_table(ws, 4, ["#", "단계", "Endpoint", "받는 주요 필드", "언리얼 동작"],
            rows, [4, 12, 36, 40, 42], mono_cols=(3, 4))
ws.sheet_view.showGridLines = False

# ---------------- 데이터구조 ----------------
ws = wb.create_sheet("데이터구조")
banner(ws, "데이터 구조 — Painting 모델 + 응답 필드", "필드명은 SSOT(API 명세서)와 동일")
r = 4
r = section(ws, r, "Painting (명화 데이터 모델)")
painting = [
    ("paintingId", "string", "Y", "명화 고유 ID (예: mona_lisa)"),
    ("title", "string", "Y", "작품명 (예: 모나리자)"),
    ("artist", "string", "Y", "작가"),
    ("era", "string", "N", "시대/사조"),
    ("personaPrompt", "string", "Y", "AI 페르소나 프롬프트(서버 전용, 클라엔 보통 미노출)"),
    ("assetKey", "string", "N", "언리얼 에셋 매핑 키 — 복제 시 교체 지점"),
]
r = style_table(ws, r, ["필드", "타입", "필수", "설명"], painting, [16, 10, 6, 60], mono_cols=(1,))
r += 1
r = section(ws, r, "Chat 응답 (POST .../chat)")
chat = [
    ("reply", "string", "Y", "명화 페르소나 답변 텍스트"),
    ("ttsUrl", "string", "Y", "답변 음성 파일 URL"),
    ("turnCount", "integer", "Y", "누적 대화 턴 수"),
    ("quizReady", "boolean", "Y", "퀴즈 생성 가능 여부(예: 2턴 이상)"),
]
r = style_table(ws, r, ["필드", "타입", "필수", "설명"], chat, [16, 10, 6, 60], mono_cols=(1,))
r += 1
r = section(ws, r, "Quiz 응답 (POST .../quiz)")
quiz = [
    ("quizId", "string", "Y", "퀴즈 식별자(봉인 시 사용)"),
    ("question", "string", "Y", "문제 텍스트"),
    ("choices", "array<string>", "Y", "보기 목록(2~4개)"),
    ("answerIndex", "integer", "—", "정답 인덱스: 서버 보관, 클라엔 안 옴(판정은 /seal)"),
]
r = style_table(ws, r, ["필드", "타입", "필수", "설명"], quiz, [16, 14, 6, 56], mono_cols=(1,))
r += 1
r = section(ws, r, "Seal 응답 (POST .../seal)")
seal = [
    ("correct", "boolean", "Y", "정답 여부"),
    ("sealed", "boolean", "Y", "봉인 성공 여부(정답 시 true)"),
    ("stampCount", "integer", "Y", "현재 모은 스탬프 수"),
    ("totalStamps", "integer", "Y", "전체 목표 스탬프 수"),
    ("feedback", "string", "N", "정답 칭찬 / 오답 격려·힌트"),
]
r = style_table(ws, r, ["필드", "타입", "필수", "설명"], seal, [16, 10, 6, 60], mono_cols=(1,))
ws.sheet_view.showGridLines = False

# ---------------- USTRUCT매핑 ----------------
ws = wb.create_sheet("USTRUCT매핑")
banner(ws, "USTRUCT 매핑 — JSON ↔ 언리얼 C++", "FJsonObjectConverter 사용. UE는 PascalCase 프로퍼티↔camelCase JSON 자동 매칭")
r = 4
r = section(ws, r, "주의사항")
warn = [
    "UPROPERTY 프로퍼티명은 PascalCase로, JSON 키는 camelCase로 두면 자동 매칭됩니다 (PaintingId ↔ paintingId).",
    "bool 프로퍼티(b접두)는 자동 매칭이 어긋날 수 있어 변환을 1회 검증하세요. 안 되면 'bool Correct;' 식으로 직접 두거나 커스텀 파싱.",
    "C++ vs Blueprint 구조체 선택은 팀 미확정(CLAUDE.md §11). 아래는 C++ 예시.",
]
for n in warn:
    ws.cell(r, 1, "• " + n).font = cell_font; ws.merge_cells(f"A{r}:F{r}"); ws.cell(r,1).alignment = wrap; ws.row_dimensions[r].height = 30; r += 1
r += 1
maps = [
    ("FPaintingData", "paintingId",  "FString PaintingId", "string"),
    ("",              "title",       "FString Title",      "string"),
    ("",              "artist",      "FString Artist",     "string"),
    ("",              "era",         "FString Era",        "string"),
    ("",              "assetKey",    "FString AssetKey",   "string"),
    ("FChatResponse", "reply",       "FString Reply",      "string"),
    ("",              "ttsUrl",      "FString TtsUrl",     "string"),
    ("",              "turnCount",   "int32 TurnCount",    "integer"),
    ("",              "quizReady",   "bool bQuizReady (검증)", "boolean"),
    ("FQuizResponse", "quizId",      "FString QuizId",     "string"),
    ("",              "question",    "FString Question",   "string"),
    ("",              "choices",     "TArray<FString> Choices", "array"),
    ("FSealResponse", "correct",     "bool bCorrect (검증)",   "boolean"),
    ("",              "sealed",      "bool bSealed (검증)",    "boolean"),
    ("",              "stampCount",  "int32 StampCount",   "integer"),
    ("",              "totalStamps", "int32 TotalStamps",  "integer"),
    ("",              "feedback",    "FString Feedback",   "string"),
]
style_table(ws, r, ["USTRUCT", "JSON 키", "UE 프로퍼티", "타입"], maps, [16, 16, 30, 12], mono_cols=(2, 3))
ws.sheet_view.showGridLines = False

# ---------------- 로컬파일&전환 ----------------
ws = wb.create_sheet("로컬파일&전환")
banner(ws, "로컬 JSON 파일 & 서버 전환", "지금은 로컬 파일로 개발, 나중에 스위치만 교체")
r = 4
r = section(ws, r, "로컬 JSON 파일 (docs/data/)")
files = [
    ("mona_lisa.json", "모나리자 — MVP 풀 샘플(painting + 오프라인 응답들). 이 파일을 기준으로 복제."),
    ("the_scream.json", "뭉크(예정) — mona_lisa.json 복제 후 값/assetKey만 교체."),
    ("starry_night.json", "반 고흐(예정) — 동일하게 복제."),
]
r = style_table(ws, r, ["파일", "내용"], files, [22, 70], mono_cols=(1,))
r += 1
r = section(ws, r, "로컬 파일 구조 (mona_lisa.json)")
struct = [
    ("painting", "Painting 모델 그대로 (paintingId/title/artist/era/personaPrompt/assetKey)"),
    ("_localResponses.capture", "포획 응답 모양 샘플"),
    ("_localResponses.chatSamples[]", "대화 응답 모양 샘플(턴별)"),
    ("_localResponses.quiz", "오프라인용 미리 만든 퀴즈"),
    ("_localResponses.seal", "봉인 응답 모양 샘플"),
]
r = style_table(ws, r, ["키", "설명"], struct, [30, 62], mono_cols=(1,))
r += 1
r = section(ws, r, "서버 전환 가이드")
guide = [
    "지금: LoadPaintingFromLocalJson(\"mona_lisa.json\") 로 FPaintingData 채움.",
    "나중: RequestPaintingFromServer(\"mona_lisa\") 로 같은 구조체를 채움.",
    "데이터 모양이 같으니 포획·대화·봉인 UI 로직은 변경 없음.",
]
for n in guide:
    ws.cell(r, 1, "• " + n).font = cell_font; ws.merge_cells(f"A{r}:F{r}"); ws.cell(r,1).alignment = wrap; r += 1
r += 1
r = section(ws, r, "에셋 매핑 (assetKey → 언리얼)")
asset = [
    ("mona_lisa", "MonaLisa", "/Game/Unframed/Paintings/MonaLisa/...", "초상 텍스처·머티리얼·빈 액자 액터"),
    ("the_scream", "TheScream", "/Game/Unframed/Paintings/TheScream/...", "복제 시 이 키만 바꾸면 에셋 세트 교체"),
    ("starry_night", "StarryNight", "/Game/Unframed/Paintings/StarryNight/...", "동일"),
]
style_table(ws, r, ["paintingId", "assetKey", "에셋 경로(예시)", "비고"], asset, [16, 14, 38, 30], mono_cols=(1, 2, 3))
ws.sheet_view.showGridLines = False

out = "docs/명세서/Unframed_언리얼명세서.xlsx"
wb.save(out)
print("saved:", out, "sheets:", wb.sheetnames)

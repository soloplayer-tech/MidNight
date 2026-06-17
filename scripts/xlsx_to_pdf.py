#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""xlsx → pdf 변환기 (한글 지원). 각 시트를 그리드 표로 렌더링.
사용: python3 scripts/xlsx_to_pdf.py <in.xlsx> <out.pdf>
"""
import sys
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string, range_boundaries
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import ParagraphStyle

FONT = "AppleGothic"
pdfmetrics.registerFont(TTFont(FONT, "/System/Library/Fonts/Supplemental/AppleGothic.ttf"))

NAVY = colors.HexColor("#1B2A41")
GOLD = colors.HexColor("#D4A85A")
PANEL = colors.HexColor("#2C3A52")
LIGHT = colors.HexColor("#F2EFE6")

def cell_style(size=8, color=colors.HexColor("#1A1A1A"), bold=False, align="LEFT"):
    return ParagraphStyle("c", fontName=FONT, fontSize=size, leading=size*1.35,
                          textColor=color, wordWrap="CJK", alignment={"LEFT":0,"CENTER":1}[align])

def fill_rgb(cell):
    try:
        f = cell.fill
        if f is None or f.patternType is None:
            return None
        rgb = f.fgColor.rgb
        if isinstance(rgb, str) and len(rgb) in (6, 8):
            return "#" + rgb[-6:]
    except Exception:
        pass
    return None

def convert(inp, outp):
    wb = load_workbook(inp, data_only=True)
    page_w, page_h = landscape(A4)
    margin = 12 * mm
    avail = page_w - 2 * margin
    doc = SimpleDocTemplate(outp, pagesize=landscape(A4),
                            leftMargin=margin, rightMargin=margin,
                            topMargin=10*mm, bottomMargin=10*mm,
                            title=inp.split("/")[-1])
    story = []
    title_p = cell_style(15, colors.white, True)
    sub_p = cell_style(8.5, GOLD)
    sheet_hdr = cell_style(12, NAVY, True)

    for si, ws in enumerate(wb.worksheets):
        if si > 0:
            story.append(PageBreak())
        story.append(Paragraph("● " + ws.title, sheet_hdr))
        story.append(Spacer(1, 4))

        max_row = ws.max_row
        max_col = ws.max_column
        # trim empty trailing columns
        while max_col > 1:
            if all((ws.cell(r, max_col).value in (None, "")) for r in range(1, max_row + 1)):
                max_col -= 1
            else:
                break

        # column widths from xlsx
        widths = []
        for c in range(1, max_col + 1):
            w = ws.column_dimensions[get_column_letter(c)].width or 10
            widths.append(float(w))
        tot = sum(widths) or 1
        widths = [avail * w / tot for w in widths]

        merged = list(ws.merged_cells.ranges)
        def merged_of(r, c):
            for mr in merged:
                c1, r1, c2, r2 = range_boundaries(str(mr))
                if r1 <= r <= r2 and c1 <= c <= c2:
                    return (r1, c1, r2, c2)
            return None

        data = []
        style_cmds = [
            ("FONT", (0,0), (-1,-1), FONT, 8),
            ("GRID", (0,0), (-1,-1), 0.4, colors.HexColor("#C9C2B0")),
            ("VALIGN", (0,0), (-1,-1), "TOP"),
            ("LEFTPADDING", (0,0), (-1,-1), 4),
            ("RIGHTPADDING", (0,0), (-1,-1), 4),
            ("TOPPADDING", (0,0), (-1,-1), 3),
            ("BOTTOMPADDING", (0,0), (-1,-1), 3),
        ]
        spans_done = set()
        for r in range(1, max_row + 1):
            rowvals = []
            for c in range(1, max_col + 1):
                rc = (r, c)
                m = merged_of(r, c)
                val = ws.cell(r, c).value
                if m and (r, c) != (m[0], m[1]):
                    rowvals.append("")  # covered by span anchor
                    continue
                hexfill = fill_rgb(ws.cell(r, c))
                txtcolor = colors.HexColor("#1A1A1A")
                bold = False
                if hexfill in ("#1B2A41", "#2C3A52"):
                    txtcolor = colors.white; bold = True
                    style_cmds.append(("BACKGROUND", (c-1, r-1), (c-1, r-1),
                                       NAVY if hexfill=="#1B2A41" else PANEL))
                elif hexfill == "#D4A85A":
                    txtcolor = NAVY; bold = True
                    style_cmds.append(("BACKGROUND", (c-1, r-1), (c-1, r-1), GOLD))
                elif hexfill == "#F2EFE6":
                    style_cmds.append(("BACKGROUND", (c-1, r-1), (c-1, r-1), LIGHT))
                # banner rows
                if r == 1:
                    txtcolor = colors.white; bold = True
                    style_cmds.append(("BACKGROUND", (c-1, r-1), (c-1, r-1), NAVY))
                elif r == 2:
                    txtcolor = GOLD
                    style_cmds.append(("BACKGROUND", (c-1, r-1), (c-1, r-1), NAVY))
                size = 13 if r == 1 else (8.5 if r == 2 else 8)
                p = Paragraph(str(val) if val is not None else "",
                              cell_style(size, txtcolor, bold))
                rowvals.append(p)
                if m and m not in spans_done:
                    spans_done.add(m)
                    style_cmds.append(("SPAN", (m[1]-1, m[0]-1), (m[3]-1, m[2]-1)))
            data.append(rowvals)

        if not data:
            continue
        t = Table(data, colWidths=widths, repeatRows=0)
        t.setStyle(TableStyle(style_cmds))
        story.append(t)

    doc.build(story)
    print("saved:", outp)

def resolve(path):
    import os, unicodedata
    if os.path.exists(path):
        return path
    nfc = lambda s: unicodedata.normalize("NFC", s)
    cur = "."
    for part in path.split("/"):
        if not part:
            continue
        match = None
        for f in os.listdir(cur):
            if nfc(f) == nfc(part):
                match = f
                break
        if match is None:
            raise FileNotFoundError(path)
        cur = os.path.join(cur, match)
    return cur

if __name__ == "__main__":
    convert(resolve(sys.argv[1]), sys.argv[2])
